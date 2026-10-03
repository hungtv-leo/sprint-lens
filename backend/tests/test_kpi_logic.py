from __future__ import annotations

import asyncio
import unittest
from datetime import date, datetime
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.schemas.jira import IssuesResponse, JiraAssignee, JiraIssue
from app.schemas.kpi import (
    KpiAgentPayload,
    KpiCalculateResponse,
    KpiAgentResult,
    KpiCellUpdate,
    KpiPlanItem,
    KpiRequest,
    KpiPeriod,
    KpiStats,
)
from app.services.kpi_agent import (
    commitment_rate_to_score,
    compute_developer_kpi,
    schedule_rate_to_score,
    throughput_rate_to_score,
)
from app.services.kpi_service import KpiService


class FakeJiraService:
    def __init__(self, responses: list[IssuesResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict[str, object]] = []

    async def get_issues(
        self,
        project_keys: list[str],
        sprint: str | None = "active",
        assignee: str | None = None,
        query: str | None = None,
        updated_from: str | None = None,
        updated_to: str | None = None,
    ) -> IssuesResponse:
        self.calls.append(
            {
                "project_keys": project_keys,
                "sprint": sprint,
                "assignee": assignee,
                "query": query,
                "updated_from": updated_from,
                "updated_to": updated_to,
            }
        )
        return self.responses.pop(0)


class FakeAgentRunner:
    name = "fake-agent"

    async def run(self, payload: KpiAgentPayload) -> KpiAgentResult:
        return compute_developer_kpi(payload)


class KpiLogicTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = KpiService(jira=None, agent=None)

    def make_issue(
        self,
        key: str,
        *,
        status_category: str = "done",
        status_name: str = "Done",
        assignee_name: str = "hung.tran",
        assignee_display_name: str = "Hưng",
        updated: str | None = "2026-07-09T10:00:00Z",
        resolution_date: str | None = None,
        status_category_change_date: str | None = None,
        due_date: str | None = None,
        story_points: float | None = None,
    ) -> JiraIssue:
        return JiraIssue(
            key=key,
            summary=f"Summary {key}",
            status_id="10000",
            status_name=status_name,
            status_category=status_category,
            priority="Medium",
            assignee=JiraAssignee(name=assignee_name, display_name=assignee_display_name),
            updated=updated,
            resolution_date=resolution_date,
            status_category_change_date=status_category_change_date,
            due_date=due_date,
            issue_type="Task",
            project_key="TI",
            project_name="TI",
            sprint_names=[],
            url=f"https://jira.local/browse/{key}",
            story_points=story_points,
        )

    def test_parse_plan_sheet_warns_when_due_date_out_of_month(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["D2"] = "Hạn dự kiến"
        worksheet["E2"] = "Điểm khối lượng"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        worksheet["D3"] = "2026-08-01"
        worksheet["E3"] = 3

        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        _, _, validation_issues = self.service._parse_plan_sheet(workbook, request=request)
        codes = {item.code for item in validation_issues}
        self.assertIn("expected_due_date_out_of_month", codes)

    def test_parse_flexible_due_date_formats(self) -> None:
        cases = [
            ("2026-07-17", "2026-07-17"),
            ("17/07/2026", "2026-07-17"),
            ("17-07-2026", "2026-07-17"),
            ("17/07/26", "2026-07-17"),
            ("17-07-26", "2026-07-17"),
            (date(2026, 7, 20), "2026-07-20"),
            (datetime(2026, 7, 21, 8, 30), "2026-07-21"),
        ]
        for raw_value, expected in cases:
            with self.subTest(raw_value=raw_value):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Kế hoạch Developer"
                worksheet["A2"] = "Mã Jira"
                worksheet["B2"] = "Tên công việc"
                worksheet["C2"] = "Loại kế hoạch"
                worksheet["D2"] = "Hạn dự kiến"
                worksheet["A3"] = "TI-801"
                worksheet["B3"] = "Task A"
                worksheet["C3"] = "Cam kết"
                worksheet["D3"] = raw_value

                plan_items, _, validation_issues = self.service._parse_plan_sheet(workbook)
                self.assertEqual(validation_issues, [])
                self.assertEqual(plan_items[0].expected_due_date, expected)

    def test_parse_plan_sheet_accepts_wps_like_plan_type_text(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam k?t"

        plan_items, _, validation_issues = self.service._parse_plan_sheet(workbook)
        self.assertEqual(len(validation_issues), 0)
        self.assertEqual(plan_items[0].plan_type, "committed")

    def test_compute_developer_kpi_uses_plan_due_date_and_scope_score(self) -> None:
        payload = KpiAgentPayload(
            role="developer",
            period={"type": "month", "month": "2026-07"},
            assignee="hung.tran",
            projects=["TI"],
            plan_items=[
                KpiPlanItem(
                    row_number=3,
                    issue_key="TI-801",
                    task_title="Task A",
                    plan_type="committed",
                    expected_due_date="2026-07-10",
                    scope_score=3,
                )
            ],
            issues=[
                {
                    "key": "TI-801",
                    "summary": "Task A",
                    "status_category": "done",
                    "status_name": "Done",
                    "assignee_name": "hung.tran",
                    "assignee_display_name": "Hưng",
                    "due_date": None,
                    "resolution_date": "2026-07-09T10:00:00Z",
                    "updated": "2026-07-20T10:00:00Z",
                    "story_points": None,
                }
            ],
        )

        result = compute_developer_kpi(payload)
        self.assertEqual(result.stats.schedule_source, "plan_due_date")
        self.assertEqual(result.stats.throughput_source, "scope_score")
        self.assertEqual(result.stats.schedule_rate, 1.0)
        self.assertEqual(result.stats.throughput_rate, 1.0)
        self.assertTrue(any("nguồn hạn=plan_due_date" in item for item in result.trace))
        self.assertTrue(any("resolutiondate" in item for item in result.trace))

    def test_score_bands_match_excel(self) -> None:
        self.assertEqual(commitment_rate_to_score(0.89), 8.0)
        self.assertEqual(schedule_rate_to_score(0.89), 8.5)
        self.assertEqual(throughput_rate_to_score(0.80), 8.5)
        self.assertEqual(throughput_rate_to_score(0.79), 7.0)
        self.assertEqual(schedule_rate_to_score(0.70), 7.5)
        self.assertEqual(schedule_rate_to_score(0.69), 6.0)

    def test_compute_uses_resolution_date_not_updated_for_schedule(self) -> None:
        payload = KpiAgentPayload(
            role="developer",
            period={"type": "month", "month": "2026-07"},
            assignee="hung.tran",
            projects=["TI"],
            issues=[
                {
                    "key": "TI-1",
                    "summary": "Done early then edited late",
                    "status_category": "done",
                    "status_name": "Done",
                    "due_date": "2026-07-10",
                    "resolution_date": "2026-07-09T08:00:00Z",
                    "updated": "2026-07-25T10:00:00Z",
                    "story_points": 2,
                }
            ],
        )
        result = compute_developer_kpi(payload)
        self.assertEqual(result.stats.schedule_rate, 1.0)
        self.assertEqual(result.stats.commitment_rate, 1.0)
        note = next(item.note for item in result.evidence if item.key == "TI-1")
        self.assertIn("resolutiondate", note)
        self.assertIn("Đúng hạn", note)

    def test_compute_zero_committed_returns_not_applicable(self) -> None:
        payload = KpiAgentPayload(
            role="developer",
            period={"type": "month", "month": "2026-07"},
            assignee="hung.tran",
            projects=["TI"],
            issues=[],
        )
        result = compute_developer_kpi(payload)
        self.assertIsNone(result.stats.commitment_rate)
        self.assertEqual(next(item.value for item in result.cell_updates if item.cell == "J5"), "x")
        self.assertEqual(next(item.value for item in result.cell_updates if item.cell == "K5"), "x")

    def test_throughput_allows_partial_story_points_and_scope_hybrid(self) -> None:
        payload = KpiAgentPayload(
            role="developer",
            period={"type": "month", "month": "2026-07"},
            assignee="hung.tran",
            projects=["TI"],
            plan_items=[
                KpiPlanItem(
                    row_number=3,
                    issue_key="TI-1",
                    task_title="Has SP",
                    plan_type="committed",
                    scope_score=None,
                ),
                KpiPlanItem(
                    row_number=4,
                    issue_key="TI-2",
                    task_title="Has scope only",
                    plan_type="committed",
                    scope_score=5,
                ),
                KpiPlanItem(
                    row_number=5,
                    issue_key="TI-3",
                    task_title="Missing both",
                    plan_type="committed",
                ),
            ],
            issues=[
                {
                    "key": "TI-1",
                    "summary": "Has SP",
                    "status_category": "done",
                    "status_name": "Done",
                    "due_date": "2026-07-10",
                    "resolution_date": "2026-07-09T08:00:00Z",
                    "story_points": 3,
                },
                {
                    "key": "TI-2",
                    "summary": "Has scope only",
                    "status_category": "done",
                    "status_name": "Done",
                    "due_date": None,
                    "resolution_date": "2026-07-09T08:00:00Z",
                    "story_points": None,
                },
                {
                    "key": "TI-3",
                    "summary": "Missing both",
                    "status_category": "indeterminate",
                    "status_name": "In Progress",
                    "story_points": None,
                },
            ],
        )
        result = compute_developer_kpi(payload)
        self.assertEqual(result.stats.throughput_source, "hybrid")
        self.assertEqual(result.stats.throughput_rate, 1.0)
        self.assertEqual(result.stats.throughput_committed_scope, 8.0)
        self.assertAlmostEqual(result.stats.throughput_coverage or 0, 2 / 3, places=4)
        self.assertTrue(any("Thiếu khối lượng" in note for note in result.notes))

    def test_schedule_partial_due_dates_sets_coverage(self) -> None:
        payload = KpiAgentPayload(
            role="developer",
            period={"type": "sprint", "sprint": "1"},
            assignee="hung.tran",
            projects=["TI"],
            issues=[
                {
                    "key": "TI-1",
                    "summary": "With due",
                    "status_category": "done",
                    "status_name": "Done",
                    "due_date": "2026-07-10",
                    "resolution_date": "2026-07-09T08:00:00Z",
                    "story_points": 1,
                },
                {
                    "key": "TI-2",
                    "summary": "No due",
                    "status_category": "done",
                    "status_name": "Done",
                    "due_date": None,
                    "resolution_date": "2026-07-09T08:00:00Z",
                    "story_points": 1,
                },
            ],
        )
        result = compute_developer_kpi(payload)
        self.assertEqual(result.stats.on_time_eligible, 1)
        self.assertEqual(result.stats.schedule_rate, 1.0)
        self.assertEqual(result.stats.schedule_coverage, 0.5)
        k6 = next(item.value for item in result.cell_updates if item.cell == "K6")
        self.assertEqual(k6, 10.0)

    def test_apply_cell_updates_formats_plain_number_cells(self) -> None:
        calculated = KpiCalculateResponse(
            role="developer",
            period=KpiPeriod(type="month", month="2026-07"),
            assignee="hung.tran",
            projects=["TI"],
            issue_count=1,
            result=KpiAgentResult(
                stats=KpiStats(
                    committed=1,
                    completed=1,
                    incomplete=0,
                    on_time_completed=1,
                    on_time_eligible=1,
                    commitment_rate=1.0,
                    schedule_rate=1.0,
                    throughput_rate=1.0,
                ),
                cell_updates=[
                    KpiCellUpdate(sheet="KPI Developer Demo", cell="H5", value=3),
                    KpiCellUpdate(sheet="KPI Developer Demo", cell="I5", value=3),
                    KpiCellUpdate(sheet="KPI Developer Demo", cell="J5", value=1.0),
                    KpiCellUpdate(sheet="KPI Developer Demo", cell="K5", value=10),
                ],
                trace=["Cam kết: 1/1 issue hoàn thành."],
            ),
            agent="local-skill",
        )

        content = self.service.apply_cell_updates(calculated)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(
            set(workbook.sheetnames),
            {
                "KPI Developer Demo",
                "Kế hoạch Developer",
                "Log tính KPI",
            },
        )
        worksheet = workbook["KPI Developer Demo"]
        self.assertEqual(worksheet["H5"].number_format, "0.##")
        self.assertEqual(worksheet["H5"].value, 3)
        self.assertEqual(worksheet["J5"].value, 1.0)
        self.assertEqual(worksheet["K5"].value, 10)
        # CBQL starts equal to employee scores (manager can adjust later).
        self.assertEqual(worksheet["L5"].value, worksheet["H5"].value)
        self.assertEqual(worksheet["M5"].value, worksheet["I5"].value)
        self.assertEqual(worksheet["N5"].value, worksheet["J5"].value)
        self.assertEqual(worksheet["O5"].value, worksheet["K5"].value)
        self.assertEqual(worksheet["O8"].value, worksheet["K8"].value)
        self.assertTrue(str(worksheet["O4"].value or "").startswith("="))
        log_sheet = workbook["Log tính KPI"]
        self.assertEqual(log_sheet["A1"].value, "Thông tin chung")

    def test_lead_developer_writes_and_exports_lead_sheet(self) -> None:
        payload = KpiAgentPayload(
            role="lead_developer",
            period={"type": "month", "month": "2026-07"},
            assignee="hung.tran",
            projects=["TI"],
            plan_items=[
                KpiPlanItem(
                    row_number=3,
                    issue_key="TI-1",
                    task_title="Task",
                    plan_type="committed",
                    expected_due_date="2026-07-20",
                    scope_score=3,
                )
            ],
            issues=[
                {
                    "key": "TI-1",
                    "summary": "Task",
                    "status_category": "done",
                    "status_name": "Done",
                    "assignee_name": "hung.tran",
                    "assignee_display_name": "Hưng",
                    "due_date": "2026-07-20",
                    "resolution_date": "2026-07-15T10:00:00Z",
                    "updated": "2026-07-15T10:00:00Z",
                    "story_points": 3,
                }
            ],
        )
        result = compute_developer_kpi(payload)
        self.assertTrue(all(item.sheet == "KPI Lead Developer Demo" for item in result.cell_updates))

        calculated = KpiCalculateResponse(
            role="lead_developer",
            period=KpiPeriod(type="month", month="2026-07"),
            assignee="hung.tran",
            projects=["TI"],
            issue_count=1,
            result=result,
            agent="local-skill",
        )
        content = self.service.apply_cell_updates(calculated)
        workbook = load_workbook(BytesIO(content))
        self.assertIn("KPI Lead Developer Demo", workbook.sheetnames)
        self.assertNotIn("KPI Developer Demo", workbook.sheetnames)
        lead = workbook["KPI Lead Developer Demo"]
        self.assertEqual(lead["J5"].value, result.stats.commitment_rate)
        self.assertEqual(lead["N5"].value, lead["J5"].value)
        self.assertEqual(lead["K18"].value, 0)
        self.assertEqual(lead["O18"].value, 0)
        self.assertEqual(lead["K23"].value, 0)
        self.assertTrue("0.1*K18" in str(lead["K3"].value))
        self.assertTrue("0.2*K23" in str(lead["K3"].value))

    def test_fetch_issues_requires_specific_sprint(self) -> None:
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="sprint", sprint=None),
        )

        with self.assertRaises(HTTPException) as exc:
            asyncio.run(self.service._fetch_issues_response(request))
        self.assertEqual(exc.exception.status_code, 400)
        self.assertIn("yêu cầu chọn một sprint cụ thể", exc.exception.detail)

    def test_calculate_marks_truncated_jira_as_blocking(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        buffer = BytesIO()
        workbook.save(buffer)

        jira = FakeJiraService(
            [
                IssuesResponse(
                    issues=[self.make_issue("TI-801")],
                    total=1200,
                    returned=1,
                    truncated=True,
                ),
                IssuesResponse(
                    issues=[self.make_issue("TI-801")],
                    total=1,
                    returned=1,
                    truncated=False,
                ),
            ]
        )
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        calculated = asyncio.run(service.calculate(request, workbook_bytes=buffer.getvalue()))
        truncated_issues = [item for item in calculated.validation_issues if item.code == "jira_truncated"]
        self.assertTrue(calculated.dataset_truncated)
        self.assertEqual(len(truncated_issues), 1)
        self.assertTrue(truncated_issues[0].blocking)

    def test_calculate_treats_done_issue_outside_month_as_incomplete_for_month_kpi(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["D2"] = "Hạn dự kiến"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        worksheet["D3"] = "2026-07-10"
        buffer = BytesIO()
        workbook.save(buffer)

        jira = FakeJiraService(
            [
                IssuesResponse(
                    issues=[
                        self.make_issue(
                            "TI-801",
                            updated="2026-07-20T10:00:00Z",
                            resolution_date="2026-06-30T10:00:00Z",
                        )
                    ],
                    total=1,
                    returned=1,
                    truncated=False,
                ),
                IssuesResponse(
                    issues=[],
                    total=0,
                    returned=0,
                    truncated=False,
                ),
            ]
        )
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        calculated = asyncio.run(service.calculate(request, workbook_bytes=buffer.getvalue()))
        self.assertEqual(calculated.issue_count, 1)
        self.assertEqual(calculated.result.stats.completed, 0)
        self.assertEqual(calculated.result.stats.incomplete, 1)
        self.assertTrue(
            any(item.code == "outside_requested_month" for item in calculated.validation_issues)
        )

    def test_calculate_keeps_incomplete_planned_issue_without_month_activity_in_denominator(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        buffer = BytesIO()
        workbook.save(buffer)

        jira = FakeJiraService(
            [
                IssuesResponse(
                    issues=[
                        self.make_issue(
                            "TI-801",
                            status_category="indeterminate",
                            status_name="In Progress",
                            updated="2026-06-20T10:00:00Z",
                        )
                    ],
                    total=1,
                    returned=1,
                    truncated=False,
                ),
                IssuesResponse(
                    issues=[],
                    total=0,
                    returned=0,
                    truncated=False,
                ),
            ]
        )
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        calculated = asyncio.run(service.calculate(request, workbook_bytes=buffer.getvalue()))
        self.assertEqual(calculated.issue_count, 1)
        self.assertEqual(calculated.result.stats.committed, 1)
        self.assertEqual(calculated.result.stats.incomplete, 1)
        self.assertTrue(
            any(item.code == "outside_requested_month" for item in calculated.validation_issues)
        )

    def test_calculate_rejects_empty_plan_sheet(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        buffer = BytesIO()
        workbook.save(buffer)

        jira = FakeJiraService([])
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        with self.assertRaises(HTTPException) as exc:
            asyncio.run(service.calculate(request, workbook_bytes=buffer.getvalue()))
        self.assertEqual(exc.exception.status_code, 400)
        self.assertIn("không có dòng hợp lệ", exc.exception.detail)

    def test_unplanned_issues_only_include_in_progress_non_subtasks(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        buffer = BytesIO()
        workbook.save(buffer)

        in_progress_issue = self.make_issue(
            "TI-900",
            status_category="indeterminate",
            status_name="In Progress",
        )
        done_issue = self.make_issue(
            "TI-901",
            status_category="done",
            status_name="Done",
        )

        jira = FakeJiraService(
            [
                IssuesResponse(
                    issues=[self.make_issue("TI-801")],
                    total=1,
                    returned=1,
                    truncated=False,
                ),
                IssuesResponse(
                    issues=[in_progress_issue, done_issue],
                    total=2,
                    returned=2,
                    truncated=False,
                ),
            ]
        )
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        calculated = asyncio.run(service.calculate(request, workbook_bytes=buffer.getvalue()))
        unplanned_items = [item for item in calculated.validation_issues if item.code == "unplanned_issue"]
        self.assertEqual(calculated.plan_summary.unplanned_issue_count, 1)
        self.assertEqual(len(unplanned_items), 1)
        self.assertEqual(unplanned_items[0].issue_key, "TI-900")

    def test_export_blocks_when_blocking_warnings_exist(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kế hoạch Developer"
        worksheet["A2"] = "Mã Jira"
        worksheet["B2"] = "Tên công việc"
        worksheet["C2"] = "Loại kế hoạch"
        worksheet["A3"] = "TI-801"
        worksheet["B3"] = "Task A"
        worksheet["C3"] = "Cam kết"
        buffer = BytesIO()
        workbook.save(buffer)

        jira = FakeJiraService(
            [
                IssuesResponse(
                    issues=[],
                    total=0,
                    returned=0,
                    truncated=False,
                ),
                IssuesResponse(
                    issues=[],
                    total=0,
                    returned=0,
                    truncated=False,
                ),
            ]
        )
        service = KpiService(jira=jira, agent=FakeAgentRunner())
        request = KpiRequest(
            role="developer",
            assignee="hung.tran",
            projects=["TI"],
            period=KpiPeriod(type="month", month="2026-07"),
        )

        with self.assertRaises(HTTPException) as exc:
            asyncio.run(service.export(request, workbook_bytes=buffer.getvalue()))
        self.assertEqual(exc.exception.status_code, 400)
        self.assertIn("Cần xử lý cảnh báo chặn", exc.exception.detail)


if __name__ == "__main__":
    unittest.main()
