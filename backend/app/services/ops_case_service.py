from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from fastapi import HTTPException

from app.schemas.ops import (
    OpsCaseCreate,
    OpsCaseItem,
    OpsCaseListResponse,
    OpsCaseStatsResponse,
    OpsHandlerStat,
    OpsWorkType,
    OpsWorkTypeStat,
)

WORK_TYPES: list[OpsWorkType] = [
    OpsWorkType(code="OP01_SUPPORT", label="Hỗ trợ case nghiệp vụ (CS / FAQ vượt)"),
    OpsWorkType(code="OP01_DATA", label="Kiểm tra / chỉnh dữ liệu hệ thống"),
    OpsWorkType(code="OP02_CHECK", label="Kiểm tra hệ thống định kỳ"),
    OpsWorkType(code="OP02_INCIDENT", label="Xử lý sự cố hệ thống"),
    OpsWorkType(code="OP02_BACKUP", label="Backup / monitoring / lưu trữ"),
    OpsWorkType(code="OP03_HANDOVER", label="Nghiệm thu / hướng dẫn / chuyển giao tính năng"),
    OpsWorkType(code="OP04_ACCOUNT", label="Cấp / thu hồi / phân quyền tài khoản"),
    OpsWorkType(code="OP05_EVENT", label="Vận hành kỳ thi / sự kiện"),
    OpsWorkType(code="OTHER", label="Khác (trong phạm vi vận hành)"),
]

WORK_TYPE_LABELS = {item.code: item.label for item in WORK_TYPES}


def default_db_path() -> Path:
    return Path(__file__).resolve().parents[2] / "data" / "ops_cases.db"


class OpsCaseService:
    def __init__(self, db_path: Path | None = None) -> None:
        self.db_path = db_path or default_db_path()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row
        return connection

    def _init_db(self) -> None:
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS ops_cases (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    handler TEXT NOT NULL,
                    handler_display_name TEXT NOT NULL,
                    work_type TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_ops_cases_created_at ON ops_cases(created_at)"
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_ops_cases_handler ON ops_cases(handler)"
            )
            connection.commit()

    def list_work_types(self) -> list[OpsWorkType]:
        return list(WORK_TYPES)

    def _validate_work_type(self, work_type: str) -> str:
        code = work_type.strip()
        if code not in WORK_TYPE_LABELS:
            raise HTTPException(status_code=400, detail=f"Loại công việc không hợp lệ: {work_type}")
        return code

    def _row_to_item(self, row: sqlite3.Row) -> OpsCaseItem:
        work_type = row["work_type"]
        return OpsCaseItem(
            id=row["id"],
            handler=row["handler"],
            handler_display_name=row["handler_display_name"],
            work_type=work_type,
            work_type_label=WORK_TYPE_LABELS.get(work_type, work_type),
            created_at=row["created_at"],
        )

    def create(self, payload: OpsCaseCreate) -> OpsCaseItem:
        work_type = self._validate_work_type(payload.work_type)
        handler = payload.handler.strip()
        handler_display_name = payload.handler_display_name.strip()
        if not handler or not handler_display_name:
            raise HTTPException(status_code=400, detail="Thiếu người xử lý.")

        created_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO ops_cases (handler, handler_display_name, work_type, created_at)
                VALUES (?, ?, ?, ?)
                """,
                (handler, handler_display_name, work_type, created_at),
            )
            connection.commit()
            case_id = int(cursor.lastrowid)

        return OpsCaseItem(
            id=case_id,
            handler=handler,
            handler_display_name=handler_display_name,
            work_type=work_type,
            work_type_label=WORK_TYPE_LABELS[work_type],
            created_at=created_at,
        )

    def list_cases(
        self,
        month: str | None = None,
        handler: str | None = None,
        limit: int = 100,
    ) -> OpsCaseListResponse:
        clauses: list[str] = []
        params: list[object] = []

        if month:
            if len(month) != 7 or month[4] != "-":
                raise HTTPException(status_code=400, detail="month phải có dạng YYYY-MM.")
            clauses.append("substr(created_at, 1, 7) = ?")
            params.append(month)

        if handler and handler.strip():
            clauses.append("handler = ?")
            params.append(handler.strip())

        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        safe_limit = max(1, min(limit, 10_000))

        with self._connect() as connection:
            total_row = connection.execute(
                f"SELECT COUNT(*) AS total FROM ops_cases {where}",
                params,
            ).fetchone()
            rows = connection.execute(
                f"""
                SELECT id, handler, handler_display_name, work_type, created_at
                FROM ops_cases
                {where}
                ORDER BY created_at DESC, id DESC
                LIMIT ?
                """,
                [*params, safe_limit],
            ).fetchall()

        return OpsCaseListResponse(
            items=[self._row_to_item(row) for row in rows],
            total=int(total_row["total"] if total_row else 0),
        )

    def stats(self, month: str) -> OpsCaseStatsResponse:
        if len(month) != 7 or month[4] != "-":
            raise HTTPException(status_code=400, detail="month phải có dạng YYYY-MM.")

        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT handler, handler_display_name, work_type, COUNT(*) AS total
                FROM ops_cases
                WHERE substr(created_at, 1, 7) = ?
                GROUP BY handler, handler_display_name, work_type
                ORDER BY total DESC, handler_display_name ASC
                """,
                (month,),
            ).fetchall()

        handler_map: dict[str, OpsHandlerStat] = {}
        work_type_totals: dict[str, int] = {}
        total = 0

        for row in rows:
            count = int(row["total"])
            total += count
            handler = row["handler"]
            work_type = row["work_type"]
            work_type_totals[work_type] = work_type_totals.get(work_type, 0) + count

            existing = handler_map.get(handler)
            if existing is None:
                handler_map[handler] = OpsHandlerStat(
                    handler=handler,
                    handler_display_name=row["handler_display_name"],
                    total=count,
                    by_work_type={work_type: count},
                )
            else:
                existing.total += count
                existing.by_work_type[work_type] = (
                    existing.by_work_type.get(work_type, 0) + count
                )

        by_handler = sorted(
            handler_map.values(),
            key=lambda item: (-item.total, item.handler_display_name.lower()),
        )
        by_work_type = [
            OpsWorkTypeStat(
                work_type=code,
                work_type_label=WORK_TYPE_LABELS.get(code, code),
                total=count,
            )
            for code, count in sorted(work_type_totals.items(), key=lambda item: (-item[1], item[0]))
        ]

        return OpsCaseStatsResponse(
            month=month,
            total=total,
            by_handler=by_handler,
            by_work_type=by_work_type,
        )

    def export_month_excel(self, month: str) -> bytes:
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Font

        stats = self.stats(month)
        cases = self.list_cases(month=month, limit=5000)

        workbook = Workbook()

        cases_sheet = workbook.active
        cases_sheet.title = "Danh sách case"
        cases_headers = ["STT", "Thời điểm (UTC)", "Người xử lý", "Mã loại", "Loại công việc"]
        cases_sheet.append(cases_headers)
        for col in range(1, len(cases_headers) + 1):
            cases_sheet.cell(1, col).font = Font(bold=True)

        # List is newest-first; export oldest-first for readable chronology.
        for index, item in enumerate(reversed(cases.items), start=1):
            cases_sheet.append(
                [
                    index,
                    item.created_at,
                    item.handler_display_name,
                    item.work_type,
                    item.work_type_label,
                ]
            )

        handler_sheet = workbook.create_sheet("Theo nhân viên")
        handler_sheet.append(["Nhân viên", "Số case"])
        handler_sheet.cell(1, 1).font = Font(bold=True)
        handler_sheet.cell(1, 2).font = Font(bold=True)
        for item in stats.by_handler:
            handler_sheet.append([item.handler_display_name, item.total])
        handler_sheet.append(["Tổng", stats.total])
        handler_sheet.cell(handler_sheet.max_row, 1).font = Font(bold=True)
        handler_sheet.cell(handler_sheet.max_row, 2).font = Font(bold=True)

        type_sheet = workbook.create_sheet("Theo loại công việc")
        type_sheet.append(["Mã loại", "Loại công việc", "Số case"])
        for col in range(1, 4):
            type_sheet.cell(1, col).font = Font(bold=True)
        for item in stats.by_work_type:
            type_sheet.append([item.work_type, item.work_type_label, item.total])
        type_sheet.append(["", "Tổng", stats.total])
        type_sheet.cell(type_sheet.max_row, 2).font = Font(bold=True)
        type_sheet.cell(type_sheet.max_row, 3).font = Font(bold=True)

        for sheet in workbook.worksheets:
            for column_cells in sheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))
                sheet.column_dimensions[column_letter].width = min(max_length + 2, 48)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
